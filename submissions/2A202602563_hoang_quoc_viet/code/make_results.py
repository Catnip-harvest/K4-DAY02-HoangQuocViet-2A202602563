"""make_results.py - Bước 5 (CPU): gom output của các phiên Kaggle thành sản phẩm nộp bài.

    python make_results.py --inputs <out_A> <out_B> <out_C> <out_D> --sub .. --data <thư mục có images/ và labels/>

Tạo trong thư mục bài nộp (--sub):
    results.xlsx (Backbones, Training, Inference, Final, PerClass, Latency, Summary)
    curves/      (chép ảnh đường cong của mọi run B/T/F)
    predictions/ (chép file dự đoán chung kết/mốc từ phiên D)
    figures/     (EDA, augmentation, đánh đổi độ chính xác-độ trễ, ma trận nhầm lẫn, ảnh bị đoán sai...)
    numbers.json (mọi con số dùng trong report.md, để đối chiếu)
Mọi chỉ số test được tính lại từ predictions/ bằng eval.py (cùng định nghĩa với lúc chấm).
"""
from __future__ import annotations

import argparse
import glob
import json
import shutil
import sys
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from train import REPO_ROOT  # noqa: E402

sys.path.insert(0, str(REPO_ROOT))
import eval as EV  # noqa: E402
from run_experiments import ABLATIONS, BACKBONES  # noqa: E402

CLASS = EV.CLASS_NAMES
AXIS_NAME = {"A": "A khởi tạo", "B": "B augmentation", "C": "C loss", "D": "D cân bằng mẫu", "F": "F chính quy hoá (EMA)",
             "combo": "kết hợp"}


def load_summaries(inputs):
    rows = {}
    for root in inputs:
        for f in glob.glob(str(Path(root) / "runs" / "*" / "seed*" / "summary.json")):
            s = json.loads(Path(f).read_text())
            s["run_dir"] = str(Path(f).parent)
            s["config"] = json.loads((Path(f).parent / "config.json").read_text())
            rows[(s["exp_id"], s["seed"])] = s
    return rows


def first(inputs, rel):
    for root in inputs:
        p = Path(root) / rel
        if p.exists():
            return p
    return None


def diff_desc(cfg: dict, base: dict) -> str:
    keys = ["init", "aug", "mix", "mix_alpha", "loss", "label_smoothing", "focal_gamma", "sampler", "ema_decay",
            "epochs", "lr_backbone", "lr_head", "img_size", "drop_rate"]
    d = [f"{k}={cfg.get(k)}" for k in keys if cfg.get(k) != base.get(k)]
    return ", ".join(d) if d else "(nền)"


def pm(mean, std, d=4):
    return f"{mean:.{d}f} ± {std:.{d}f}"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--inputs", nargs="+", required=True)
    ap.add_argument("--sub", default=str(HERE.parent))
    ap.add_argument("--data", required=True)
    args = ap.parse_args()
    sub, data = Path(args.sub), Path(args.data)
    inputs = [Path(p) for p in args.inputs]
    (sub / "figures").mkdir(parents=True, exist_ok=True)
    (sub / "curves").mkdir(parents=True, exist_ok=True)
    (sub / "predictions").mkdir(parents=True, exist_ok=True)
    num = {}
    S = load_summaries(inputs)
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    # ---------------- chép curves, figures, predictions ----------------
    for (e, k), s in S.items():
        src = Path(s["run_dir"]).parents[2] / "curves" / s["curve"]
        if src.exists():
            shutil.copy(src, sub / "curves" / s["curve"])
    for root in inputs:
        for f in glob.glob(str(root / "figures" / "*.png")):
            shutil.copy(f, sub / "figures" / Path(f).name)
    pred_src = first(inputs, "predictions")
    if pred_src:
        for f in pred_src.glob("*.csv"):
            shutil.copy(f, sub / "predictions" / f.name)

    # ---------------- độ trễ ----------------
    lat_rows = []
    for rel in ("inference/latency.json",):
        p = first(inputs, rel)
        if p:
            lat_rows += json.loads(p.read_text())
    rec_p = first(inputs, "final/final_record.json")
    rec = json.loads(rec_p.read_text()) if rec_p else None
    if rec:
        lat_rows += rec.get("latency", [])
    lat = pd.DataFrame(lat_rows)

    def bb_lat(name, batch=1, dtype="fp32", key="p50"):
        if lat.empty:
            return np.nan
        m = lat[(lat.config == f"backbone {name}") & (lat.batch == batch) & (lat.dtype == dtype)]
        return float(m[key].iloc[0]) if len(m) else np.nan

    # ---------------- Backbones ----------------
    rows = []
    for e, name, desc in BACKBONES:
        s = S.get((e, 0))
        if not s:
            continue
        rows.append({"exp_id": e, "backbone": name.split(".")[0], "tag trọng số (timm)": s["weight_tag"],
                     "#tham số (M)": s["params_m"], "GMAC (224)": s["gmacs"], "độ phân giải": s["img_size"],
                     "epoch": s["epochs"], "seed": s["seed"], "macro-F1 val": s["val_macro_f1"],
                     "top-1 val": s["val_top1"], "F1 val Chinee apple": s["val_f1_per_class"][0],
                     "F1 val Snake weed": s["val_f1_per_class"][7], "best epoch": s["best_epoch"],
                     "thời gian train/epoch (s)": s["train_time_per_epoch_s"],
                     "độ trễ batch-1 FP32 p50 (ms)": bb_lat(name), "độ trễ batch-1 AMP p50 (ms)": bb_lat(name, 1, "amp"),
                     "ảnh/s batch-32 AMP": bb_lat(name, 32, "amp", "images_per_s"),
                     "ghi chú": f"curves/{s['curve']}; 2 job song song trên 2 GPU T4 nên thời gian/epoch chỉ so sánh tương đối"})
    bb = pd.DataFrame(rows)

    # ---------------- Training ----------------
    t00 = [S[("T00", k)] for k in (0, 1, 2) if ("T00", k) in S]
    trows = []
    if t00:
        f1s = np.array([s["val_macro_f1"] for s in t00])
        t00_mean, t00_std = float(f1s.mean()), float(f1s.std(ddof=1)) if len(f1s) > 1 else float("nan")
        num["T00_val_f1_seeds"] = f1s.tolist()
        num["T00_val_f1_mean"], num["T00_val_f1_std"] = t00_mean, t00_std
        base_cfg = t00[0]["config"]
        axes = {e: ax for e, _, ax, _ in ABLATIONS}
        t_ids = sorted({e for (e, k) in S if e.startswith("T")})
        for e in t_ids:
            for k in sorted(kk for (ee, kk) in S if ee == e):
                s = S[(e, k)]
                delta = s["val_macro_f1"] - t00_mean
                if e == "T00":
                    verdict = "mốc (3 seed để đo nhiễu)"
                elif abs(delta) <= t00_std:
                    verdict = "không phân biệt được (|Δ| ≤ std)"
                elif abs(delta) <= 2 * t00_std:
                    verdict = ("tốt hơn" if delta > 0 else "kém hơn") + " nhưng < 2·std (1 seed): yếu"
                else:
                    verdict = ("tốt hơn" if delta > 0 else "kém hơn") + " rõ (> 2·std, 1 seed)"
                trows.append({"exp_id": e, "backbone": s["backbone"].split(".")[0],
                              "trục": AXIS_NAME.get(axes.get(e, "combo" if e != "T00" else ""), "nền"),
                              "khác T00 ở điểm nào": diff_desc(s["config"], base_cfg), "seed": k,
                              "macro-F1 val": s["val_macro_f1"], "top-1 val": s["val_top1"],
                              "Δ macro-F1 so với T00 (mean 3 seed)": delta, "std T00 (3 seed)": t00_std,
                              "F1 val Chinee apple": s["val_f1_per_class"][0],
                              "F1 val Snake weed": s["val_f1_per_class"][7],
                              "best epoch": s["best_epoch"], "kết luận": verdict, "ảnh": f"curves/{s['curve']}"})
    tr = pd.DataFrame(trows)

    # ---------------- Inference ----------------
    inf_p = first(inputs, "inference/inference_val.json")
    irows, temps = [], {}
    if inf_p:
        inf = json.loads(inf_p.read_text())
        temps = inf["temperature"]
        num["temperature_val"] = temps

        def lat_of(method, batch, key):
            if lat.empty or "method" not in lat:
                return np.nan
            m = lat[(lat.method == method) & (lat.batch == batch) & (lat.dtype == "amp")]
            m = m[~m.config.str.startswith("F01") & ~m.config.str.startswith("R01")]
            return float(m[key].iloc[0]) if len(m) else np.nan

        base_p50 = lat_of("I00_1view", 1, "p50")
        for r in inf["rows"]:
            meth = r["method"].replace("+TS", "")
            lm = meth if meth in ("I00_1view", "I01_hflip", "I02a_5crop", "I02b_10crop", "I02c_multiscale",
                                  "I04_full256", "I04_full288", "I04_full320", "ensemble_backbones") else None
            if meth == "I04_full224" or meth == "fp16_half" or r["method"].startswith("I06"):
                lm = "I00_1view"
            p50 = lat_of(lm, 1, "p50") if lm else np.nan
            irows.append({"exp_id": r["exp_id"], "phương pháp": r["method"], "gộp": r["aggregation"],
                          "mô hình/checkpoint": r["model"], "K (view hoặc model)": r["K"],
                          "macro-F1 val": r["val_macro_f1"], "top-1 val": r["val_top1"], "ECE val": r["val_ece"],
                          "NLL val": r["val_nll"], "F1 val Chinee": r["val_f1_chinee"], "F1 val Snake": r["val_f1_snake"],
                          "p50 b1 (ms)": p50, "p95 b1 (ms)": lat_of(lm, 1, "p95") if lm else np.nan,
                          "p99 b1 (ms)": lat_of(lm, 1, "p99") if lm else np.nan,
                          "ảnh/s b32": lat_of(lm, 32, "images_per_s") if lm else np.nan,
                          "chi phí tương đối (p50 / p50 I00)": p50 / base_p50 if lm else np.nan,
                          "số nhãn đổi so với I00": r.get("n_changed", np.nan),
                          "sai→đúng": r.get("wrong_to_right", np.nan), "đúng→sai": r.get("right_to_wrong", np.nan),
                          "ghi chú": r.get("note", "") + (" | độ trễ: AMP, view chạy tuần tự, không tính tiền xử lý"
                                                          if lm else " | độ trễ: xem sheet Latency")})
    inf_df = pd.DataFrame(irows)

    # ---------------- Final + PerClass (tính lại bằng eval.py từ predictions/) ----------------
    P = sub / "predictions"
    labels = data / "labels"
    fin_rows, groups = [], {}
    names = EV.load_names(str(labels / "labels.csv"))
    desc_cfg = {}
    if rec:
        desc_cfg = {"F01": f"{rec['method']} (gộp {rec['space']}) + TS", "F01_uncal": f"{rec['method']} chưa TS",
                    "R01": f"{rec['rt_method']} + TS", "T00": "công thức nền + I00 1-view"}
    for tag in ("F01", "F01_uncal", "R01", "T00"):
        files = sorted(glob.glob(str(P / f"{tag}_seed*_test.csv")))
        if not files:
            continue
        g = EV.load_group(files, str(labels / "test_subset0.csv"))
        groups[tag] = g
        vfiles = sorted(glob.glob(str(P / f"{tag}_seed*_val.csv")))
        gv = EV.load_group(vfiles, str(labels / "val_subset0.csv"), ref_what="val") if vfiles else None
        for i, (p, m) in enumerate(zip(g.preds, g.metrics)):
            vm = gv.metrics[i]["macro_f1"] if gv else np.nan
            fin_rows.append({"exp_id": tag, "cấu hình": desc_cfg.get(tag, tag), "seed": p.seed, "macro-F1 val": vm,
                             "macro-F1 test": m["macro_f1"], "top-1 test": m["top1"],
                             "balanced acc test": m["balanced_acc"], "ECE test": m["ece"],
                             "recall test Chinee apple": m["recall"][0], "recall test Snake weed": m["recall"][7],
                             "file": Path(p.path).name})
        sm = g.summary
        vmean = (gv.summary["macro_f1"] if gv else (np.nan, np.nan))
        fin_rows.append({"exp_id": tag, "cấu hình": desc_cfg.get(tag, tag), "seed": "mean ± std (3 seed, ddof=1)",
                         "macro-F1 val": pm(*vmean) if gv else "", "macro-F1 test": pm(*sm["macro_f1"]),
                         "top-1 test": pm(*sm["top1"]), "balanced acc test": pm(*sm["balanced_acc"]),
                         "ECE test": pm(*sm["ece"]), "recall test Chinee apple": pm(sm["recall"][0][0], sm["recall"][1][0]),
                         "recall test Snake weed": pm(sm["recall"][0][7], sm["recall"][1][7]), "file": f"{tag}_seed*_test.csv"})
        num[f"{tag}_test"] = {k: list(map(float, sm[k])) for k in ("macro_f1", "top1", "balanced_acc", "ece", "nll")}
        num[f"{tag}_test"]["recall_mean"] = sm["recall"][0].tolist()
        num[f"{tag}_test"]["recall_std"] = sm["recall"][1].tolist()
        num[f"{tag}_test"]["precision_mean"] = sm["precision"][0].tolist()
        num[f"{tag}_test"]["f1_mean"] = sm["f1"][0].tolist()
        num[f"{tag}_test"]["f1_std"] = sm["f1"][1].tolist()
        num[f"{tag}_test"]["per_seed_macro_f1"] = [m["macro_f1"] for m in g.metrics]
        num[f"{tag}_test"]["per_seed_top1"] = [m["top1"] for m in g.metrics]
        if gv:
            num[f"{tag}_val"] = {k: list(map(float, gv.summary[k])) for k in ("macro_f1", "top1", "ece")}
    fin = pd.DataFrame(fin_rows)
    pcl = []
    for tag in ("F01", "T00", "R01"):
        if tag not in groups:
            continue
        g = groups[tag]
        for i, c in enumerate(names):
            pcl.append({"cấu hình": tag, "lớp": c, "số ảnh test": int(g.metrics[0]["support"][i]),
                        "precision": g.summary["precision"][0][i], "precision std": g.summary["precision"][1][i],
                        "recall": g.summary["recall"][0][i], "recall std": g.summary["recall"][1][i],
                        "F1": g.summary["f1"][0][i], "F1 std": g.summary["f1"][1][i]})
    pc_df = pd.DataFrame(pcl)

    # ---------------- Latency ----------------
    lat_df = pd.DataFrame()
    if not lat.empty:
        lat_df = pd.DataFrame({"cấu hình": lat.config, "GPU": lat.gpu, "dtype": lat.dtype, "batch": lat.batch,
                               "độ phân giải đầu vào": lat.img_size, "K view/model": lat.k_views,
                               "gộp BN": lat.fused_bn.map({True: "có", False: "không"}), "p50 (ms)": lat.p50,
                               "p95 (ms)": lat.p95, "p99 (ms)": lat.p99, "mean (ms)": lat["mean"],
                               "ảnh/s": lat.images_per_s, "số lần đo": lat.n, "warmup": lat.warmup,
                               "torch": lat.torch, "tính tiền xử lý": "không"})

    # ---------------- Summary: top-10 theo macro-F1 val ----------------
    cands = []
    for _, r in bb.iterrows():
        cands.append({"cấu hình": f"{r['exp_id']} {r['backbone']} (nền, 1-view)", "nhóm": "Backbone", "seed": "0",
                      "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"],
                      "p50 b1 (ms)": r["độ trễ batch-1 AMP p50 (ms)"], "GMAC": r["GMAC (224)"],
                      "thời gian train/epoch (s)": r["thời gian train/epoch (s)"]})
    if len(tr):
        best_lat = bb_lat(t00[0]["backbone"], 1, "amp") if t00 else np.nan
        for _, r in tr[tr.exp_id != "T00"].iterrows():
            cands.append({"cấu hình": f"{r['exp_id']} {r['khác T00 ở điểm nào']}", "nhóm": "Training", "seed": "0",
                          "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"], "p50 b1 (ms)": best_lat,
                          "GMAC": bb[bb.backbone == r["backbone"]]["GMAC (224)"].iloc[0] if len(bb) else np.nan,
                          "thời gian train/epoch (s)": S[(r["exp_id"], 0)]["train_time_per_epoch_s"]})
        cands.append({"cấu hình": "T00 công thức nền (MỐC)", "nhóm": "Training", "seed": "0,1,2 (mean)",
                      "macro-F1 val": t00_mean, "top-1 val": float(np.mean([s["val_top1"] for s in t00])),
                      "p50 b1 (ms)": best_lat, "GMAC": np.nan,
                      "thời gian train/epoch (s)": float(np.mean([s["train_time_per_epoch_s"] for s in t00]))})
    if len(inf_df):
        for _, r in inf_df[inf_df.exp_id.isin(["I01", "I02a", "I02b", "I02c", "I03", "I04", "I05"])].iterrows():
            cands.append({"cấu hình": f"{r['exp_id']} {r['phương pháp']} ({r['gộp']}) trên {r['mô hình/checkpoint']}",
                          "nhóm": "Inference", "seed": "0", "macro-F1 val": r["macro-F1 val"], "top-1 val": r["top-1 val"],
                          "p50 b1 (ms)": r["p50 b1 (ms)"], "GMAC": np.nan, "thời gian train/epoch (s)": np.nan})
    if "F01" in groups and "F01_val" in num:
        f01_p50 = [x["p50"] for x in (rec or {}).get("latency", [])
                   if x["method"] == rec["method"] and x["batch"] == 1 and x["dtype"] == "amp"]
        cands.append({"cấu hình": f"F01 CHUNG KẾT ({desc_cfg['F01']})", "nhóm": "Final", "seed": "0,1,2 (mean)",
                      "macro-F1 val": num["F01_val"]["macro_f1"][0], "top-1 val": num["F01_val"]["top1"][0],
                      "p50 b1 (ms)": f01_p50[0] if f01_p50 else np.nan, "GMAC": np.nan,
                      "thời gian train/epoch (s)": float(np.mean([S[("F01", k)]["train_time_per_epoch_s"]
                                                                   for k in (0, 1, 2) if ("F01", k) in S]))})
    ranked = pd.DataFrame(cands).sort_values("macro-F1 val", ascending=False).reset_index(drop=True)
    ranked.insert(0, "hạng", range(1, len(ranked) + 1))
    summ = ranked.head(10)
    must = ranked[ranked["cấu hình"].str.contains("CHUNG KẾT|MỐC") & ~ranked.index.isin(summ.index)]
    summ = pd.concat([summ, must]).reset_index(drop=True)   # luôn kèm dòng chung kết và mốc (hạng thật giữ nguyên)
    if t00:
        summ["Δ so với mốc T00 (mean)"] = summ["macro-F1 val"] - t00_mean
    summary_note = pd.DataFrame({"ghi chú": [
        "Top-10 cấu hình theo macro-F1 VAL (chỉ số chọn cấu hình). Mốc T00 = mean 3 seed; các dòng khác 1 seed trừ F01.",
        "Độ trễ: GPU Tesla T4, AMP FP16, batch 1, p50, chỉ forward (không tính tiền xử lý). Chi tiết ở sheet Latency.",
        (f"Test (chạy 1 lần/seed): F01 macro-F1 {pm(*num['F01_test']['macro_f1'])} vs mốc T00 {pm(*num['T00_test']['macro_f1'])}"
         if "F01_test" in num and "T00_test" in num else "Test: chưa có"),
    ]})

    # ---------------- ghi xlsx + định dạng ----------------
    xlsx = sub / "results.xlsx"
    sheets = {"Backbones": bb, "Training": tr, "Inference": inf_df, "Final": fin, "PerClass": pc_df,
              "Latency": lat_df, "Summary": summ}
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        for name, df in sheets.items():
            df.to_excel(w, sheet_name=name, index=False)
        summary_note.to_excel(w, sheet_name="Summary", index=False, startrow=len(summ) + 2)
    format_xlsx(xlsx, sheets)
    num["xlsx_rows"] = {k: len(v) for k, v in sheets.items()}

    # ---------------- biểu đồ tổng hợp ----------------
    make_figures(sub / "figures", bb, tr, inf_df, lat, groups, names, data, num, rec, plt)
    (sub / "numbers.json").write_text(json.dumps(num, indent=1, ensure_ascii=False, default=float), encoding="utf-8")
    print("xong:", xlsx)


def format_xlsx(path, sheets):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path)
    best_fill = PatternFill("solid", fgColor="FFF2CC")
    base_fill = PatternFill("solid", fgColor="DDEBF7")
    head_fill = PatternFill("solid", fgColor="D9D9D9")
    key_col = {"Backbones": "macro-F1 val", "Training": "macro-F1 val", "Inference": "macro-F1 val",
               "Summary": "macro-F1 val", "Latency": None, "Final": None, "PerClass": None}
    for name, df in sheets.items():
        ws = wb[name]
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
        for j, col in enumerate(df.columns, start=1):
            letter = get_column_letter(j)
            width = max([len(str(col))] + [len(f"{v:.4f}") if isinstance(v, float) else len(str(v))
                                           for v in df[col].head(60)])
            ws.column_dimensions[letter].width = min(60, max(9, width + 2))
            if pd.api.types.is_float_dtype(df[col]):
                for row in ws.iter_rows(min_row=2, max_row=len(df) + 1, min_col=j, max_col=j):
                    for c in row:
                        c.number_format = "0.0000" if "ms" not in col and "ảnh/s" not in col and "(s)" not in col \
                            else "0.00"
        kc = key_col.get(name)
        if kc and kc in df.columns and len(df):
            best = int(pd.to_numeric(df[kc], errors="coerce").idxmax()) + 2
            for c in ws[best]:
                c.fill = best_fill
                c.font = Font(bold=True)
        if name in ("Final",):
            for r in range(2, len(df) + 2):
                if str(ws.cell(r, 3).value).startswith("mean"):
                    for c in ws[r]:
                        c.font = Font(bold=True)
                        c.fill = best_fill if ws.cell(r, 1).value == "F01" else base_fill
        if name == "Summary":
            for r in range(2, len(df) + 2):
                if "MỐC" in str(ws.cell(r, 2).value):
                    for c in ws[r]:
                        c.fill = base_fill
                if "CHUNG KẾT" in str(ws.cell(r, 2).value):
                    for c in ws[r]:
                        c.fill = best_fill
                        c.font = Font(bold=True)
    wb.save(path)


def make_figures(fig_dir, bb, tr, inf_df, lat, groups, names, data, num, rec, plt):
    # 1) backbone: macro-F1 val theo độ trễ và GMAC
    if len(bb):
        fig, ax = plt.subplots(1, 2, figsize=(12, 4.5))
        for _, r in bb.iterrows():
            for a, xcol in zip(ax, ["độ trễ batch-1 AMP p50 (ms)", "GMAC (224)"]):
                a.scatter(r[xcol], r["macro-F1 val"], s=40 + 4 * r["#tham số (M)"])
                a.annotate(f"{r['exp_id']} {r['backbone']}", (r[xcol], r["macro-F1 val"]), fontsize=8,
                           xytext=(4, 4), textcoords="offset points")
        ax[0].set_xlabel("độ trễ batch 1, AMP, p50 (ms) - T4"); ax[1].set_xlabel("GMAC (ảnh 224)")
        for a in ax:
            a.set_ylabel("macro-F1 val"); a.grid(alpha=.3)
        fig.suptitle("Bước 1: macro-F1 val theo độ trễ và GMAC (kích thước điểm ~ số tham số)")
        fig.tight_layout(); fig.savefig(fig_dir / "backbones_f1_vs_cost.png", dpi=120); plt.close(fig)
    # 2) ablation: Δ so với T00 kèm dải ±std, ±2std
    if len(tr):
        t = tr[tr.exp_id != "T00"]
        std = float(tr["std T00 (3 seed)"].iloc[0])
        fig, ax = plt.subplots(figsize=(10, 4.8))
        ax.axhspan(-std, std, color="gray", alpha=.25, label="±1 std T00 (3 seed)")
        ax.axhspan(-2 * std, 2 * std, color="gray", alpha=.1, label="±2 std")
        ax.bar(t.exp_id + "\n" + t["trục"].str.split(" ").str[0], t["Δ macro-F1 so với T00 (mean 3 seed)"],
               color=["tab:green" if d > 0 else "tab:red" for d in t["Δ macro-F1 so với T00 (mean 3 seed)"]])
        vals = t["Δ macro-F1 so với T00 (mean 3 seed)"].to_numpy()
        lo = -max(0.015, 4 * std)
        ax.set_ylim(lo, max(0.02, vals.max() * 1.3))
        for i_, v in enumerate(vals):   # cột bị cắt (trục y phóng to quanh 0) ghi giá trị thật
            ax.text(i_, max(v, lo * 0.97) + 0.0005, f"{v:+.4f}", ha="center", fontsize=7,
                    va="bottom", color="k")
        ax.axhline(0, color="k", lw=.8); ax.set_ylabel("Δ macro-F1 val so với T00 (trục cắt ở dưới)")
        ax.set_title("Bước 2: ablation công thức huấn luyện (1 seed mỗi dòng) so với nhiễu seed của T00")
        ax.legend(); ax.grid(alpha=.3, axis="y")
        fig.tight_layout(); fig.savefig(fig_dir / "ablation_delta.png", dpi=120); plt.close(fig)
    # 3) đánh đổi độ chính xác - độ trễ (Bước 3)
    if len(inf_df):
        d = inf_df.dropna(subset=["p50 b1 (ms)"])
        d = d[d["gộp"].isin(["logit", "-"]) | (d.exp_id == "I05")]
        fig, ax = plt.subplots(figsize=(9, 5.5))
        ax.scatter(d["p50 b1 (ms)"], d["macro-F1 val"], c="tab:blue")
        for _, r in d.iterrows():
            ax.annotate(r["phương pháp"], (r["p50 b1 (ms)"], r["macro-F1 val"]), fontsize=7, xytext=(3, 3),
                        textcoords="offset points")
        ax.set_xscale("log"); ax.set_xlabel("độ trễ batch 1, p50 (ms, log) - T4, AMP")
        ax.set_ylabel("macro-F1 val"); ax.grid(alpha=.3, which="both")
        ax.set_title("Bước 3: đánh đổi độ chính xác và độ trễ của phương pháp suy luận")
        fig.tight_layout(); fig.savefig(fig_dir / "inference_tradeoff.png", dpi=120); plt.close(fig)
    # 4) ma trận nhầm lẫn test (tổng 3 seed) cho F01 và T00
    for tag in ("F01", "T00"):
        if tag not in groups:
            continue
        cm = sum(m["confusion"] for m in groups[tag].metrics)
        cmn = cm / cm.sum(1, keepdims=True)
        fig, ax = plt.subplots(figsize=(8.5, 7.5))
        im = ax.imshow(cmn, cmap="Blues", vmin=0, vmax=1)
        for i in range(9):
            for j in range(9):
                ax.text(j, i, f"{cm[i, j]}\n{cmn[i, j]:.1%}" if cm[i, j] else "0", ha="center", va="center",
                        fontsize=6.5, color="white" if cmn[i, j] > .5 else "black")
        ax.set_xticks(range(9)); ax.set_xticklabels(names, rotation=45, ha="right", fontsize=8)
        ax.set_yticks(range(9)); ax.set_yticklabels(names, fontsize=8)
        ax.set_xlabel("nhãn dự đoán"); ax.set_ylabel("nhãn thật")
        ax.set_title(f"{tag}: ma trận nhầm lẫn TEST (cộng 3 seed; số ảnh và % theo hàng)")
        fig.colorbar(im, ax=ax, fraction=.046)
        fig.tight_layout(); fig.savefig(fig_dir / f"confusion_test_{tag}.png", dpi=120); plt.close(fig)
        off = cm.copy(); np.fill_diagonal(off, 0)
        i, j = np.unravel_index(off.argmax(), off.shape)
        pairs = sorted(((int(off[a, b]), names[a], names[b]) for a in range(9) for b in range(9) if off[a, b]),
                       reverse=True)[:6]
        num[f"{tag}_confusion_sum"] = cm.tolist()
        num[f"{tag}_most_confused"] = pairs
    # 5) ảnh test bị F01 (seed 0) đoán sai giữa Chinee apple và Snake weed
    if "F01" in groups:
        from PIL import Image
        p0 = groups["F01"].preds[0]
        mis = [(f, t, pr, p0.probs[i].max()) for i, (f, t, pr) in enumerate(zip(p0.filenames, p0.y_true, p0.y_pred))
               if t != pr]
        hard = [m for m in mis if {m[1], m[2]} == {0, 7}]
        other = [m for m in mis if {m[1], m[2]} != {0, 7}]
        pick = hard[:8] + other[:8]
        num["F01_seed0_n_mis"] = len(mis)
        num["F01_seed0_chinee_snake_mis"] = len(hard)
        if pick:
            cols = 8
            rows_ = int(np.ceil(len(pick) / cols))
            fig, axes = plt.subplots(rows_, cols, figsize=(2.4 * cols, 2.8 * rows_))
            for ax, (f, t, pr, conf) in zip(np.ravel(axes), pick):
                ax.imshow(Image.open(data / "images" / f))
                ax.set_title(f"thật: {names[t]}\nđoán: {names[pr]} ({conf:.2f})", fontsize=7)
            for ax in np.ravel(axes):
                ax.axis("off")
            fig.suptitle("F01 seed 0 - ảnh test bị đoán sai (hàng 1: cặp Chinee apple ↔ Snake weed; hàng 2: lỗi khác)")
            fig.tight_layout(); fig.savefig(fig_dir / "misclassified_test.png", dpi=100); plt.close(fig)


if __name__ == "__main__":
    main()
